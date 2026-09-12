from __future__ import annotations

import csv
import io
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import AsyncMock, patch

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from src.video_workflow.domain import Asset, AssetRole, AssetType, Delivery, JobStatus, JobType, ProjectBrief, RenderJob, SceneProfile, ScriptDurationAssessment, Shot, ShotContinuityMode, StyleProfile, VoiceEvent
from src.video_workflow.server.app import app
from src.video_workflow.server.routers import projects as router
from src.video_workflow.services.finalize import Finalizer
from src.video_workflow.services.projects import H3_DIRECTOR_VERSION, KeyframeBusyError, ProjectService
from src.video_workflow.services.render_queue import RenderQueue
from src.video_workflow.storage import ProjectStore


class ProjectApiTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        self.old = (router.store, router.project_service, router.render_queue, router.finalizer)
        router.store = ProjectStore(root / "api.sqlite3")
        router.project_service = ProjectService(router.store)
        router.render_queue = RenderQueue(router.store, router.project_service)
        router.finalizer = Finalizer(router.store, router.project_service)
        router.project_service.project_dir = lambda project_id: root / "projects" / project_id  # type: ignore[method-assign]
        self.client_context = TestClient(app)
        self.client = self.client_context.__enter__()

    def tearDown(self) -> None:
        self.client_context.__exit__(None, None, None)
        router.store, router.project_service, router.render_queue, router.finalizer = self.old
        self.temp.cleanup()

    def test_series_routes_create_template_and_episode(self) -> None:
        source = self.client.post("/api/projects", json={
            "title": "张小差科普",
            "story": "第一集",
            "visual_style": "统一二维插画",
        })
        self.assertEqual(source.status_code, 200, source.text)
        source_id = source.json()["id"]
        saved = self.client.post(
            f"/api/projects/{source_id}/series",
            json={"name": "张小差安全系列", "description": "社区安全科普"},
        )
        self.assertEqual(saved.status_code, 200, saved.text)
        series_id = saved.json()["series"]["id"]
        listed = self.client.get("/api/projects/series")
        self.assertEqual(listed.status_code, 200, listed.text)
        self.assertEqual(listed.json()[0]["id"], series_id)
        episode = self.client.post(
            f"/api/projects/series/{series_id}/episodes",
            json={
                "brief": {"title": "电动车入户充电", "story": "张小差劝阻王大爷"},
                "character_ids": [],
            },
        )
        self.assertEqual(episode.status_code, 200, episode.text)
        self.assertEqual(episode.json()["series_id"], series_id)
        self.assertEqual(episode.json()["episode_number"], 2)
        self.assertTrue(episode.json()["brief"]["visual_style"].startswith("统一二维插画"))
        self.assertIn("禁止真人照片", episode.json()["brief"]["visual_style"])

    def test_project_quote_is_optional_and_persisted(self) -> None:
        quoted = self.client.post(
            "/api/projects",
            json={"title": "报价项目", "story": "故事", "project_quote": 6800.5},
        )
        self.assertEqual(quoted.status_code, 200, quoted.text)
        self.assertEqual(quoted.json()["brief"]["project_quote"], 6800.5)
        unquoted = self.client.post(
            "/api/projects",
            json={"title": "未报价项目", "story": "故事"},
        )
        self.assertEqual(unquoted.status_code, 200, unquoted.text)
        self.assertIsNone(unquoted.json()["brief"]["project_quote"])

    def test_script_duration_assessment_route(self) -> None:
        project = self.client.post(
            "/api/projects",
            json={"title": "自然时长接口", "story": "人物发现危险并处理。", "target_duration_seconds": 60},
        ).json()
        result = ScriptDurationAssessment(
            natural_duration_seconds=42,
            natural_duration_min_seconds=38,
            natural_duration_max_seconds=46,
            target_duration_seconds=60,
            difference_seconds=-18,
            recommendation="expand",
            dialogue_and_narration_seconds=22,
            visual_only_seconds=18,
            transition_seconds=2,
            summary="内容不足",
        )
        with patch.object(router.project_service, "assess_story_duration", new=AsyncMock(return_value=result)):
            response = self.client.post(f"/api/projects/{project['id']}/brief/duration-assess", json={})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["natural_duration_seconds"], 42)
        self.assertEqual(response.json()["recommendation"], "expand")

    def test_scene_editor_preview_patch_and_conflict(self) -> None:
        project = router.project_service.create_project(ProjectBrief(title="场景编辑", story="奶奶在客厅"))
        profile = SceneProfile(name="客厅", description="普通家居灯光")
        project.scene_profiles = [profile]
        router.store.save_project(project)
        base = f"/api/projects/{project.id}/scene-profiles/{profile.id}"
        preview = self.client.get(base + "/reference-prompt")
        self.assertEqual(preview.status_code, 200)
        self.assertIn("普通家居灯光", preview.json()["prompt"])
        self.assertIn("绝对不出现任何人物", preview.json()["prompt"])
        payload = {"expected_version": 1, "name": "奶奶家", "description": "窗光", "continuity_notes": "保留桌椅", "reference_prompt": "EXACT USER PROMPT"}
        preview = self.client.post(base + "/reference-prompt", json=payload)
        self.assertEqual(preview.status_code, 200)
        self.assertIn("窗光", preview.json()["prompt"])
        self.assertEqual(router.store.get_project(project.id).scene_profiles[0].name, "客厅")
        saved = self.client.patch(base, json=payload)
        self.assertEqual(saved.status_code, 200, saved.text)
        self.assertEqual(self.client.get(base + "/reference-prompt").json()["prompt"], "EXACT USER PROMPT")
        self.assertEqual(self.client.patch(base, json=payload).status_code, 409)
        self.assertEqual(self.client.patch(base, json={**payload, "expected_version": 2, "name": " "}).status_code, 400)
        self.assertEqual(self.client.post(base + "/reference", json={"prompt": "old", "expected_version": 1}).status_code, 409)
        self.assertEqual(self.client.get(f"/api/projects/wrong/scene-profiles/{profile.id}/reference-prompt").status_code, 404)
        with patch.object(router.project_service, "generate_scene_reference", new=AsyncMock(side_effect=RuntimeError("服务商已出图，下载失败"))):
            failed = self.client.post(base + "/reference", json={"prompt": "custom", "expected_version": 2})
            self.assertEqual(failed.status_code, 502)
            self.assertIn("下载失败", failed.json()["detail"])
        self.assertEqual(self.client.post(base + "/reference/retry-download").status_code, 400)

    def test_scene_reference_revise_route_validates_source_and_forwards_feedback(self) -> None:
        project = router.project_service.create_project(ProjectBrief(title="场景改图", story="无人空镜"))
        profile = SceneProfile(name="大厅", description="水磨石地面")
        project.scene_profiles = [profile]
        router.store.save_project(project)
        base = f"/api/projects/{project.id}/scene-profiles/{profile.id}/reference/revise"
        self.assertEqual(self.client.post(base, json={
            "source_asset_id": "", "user_suggestions": "改为清晨", "expected_version": 1,
        }).status_code, 422)

        mocked = AsyncMock(return_value=Asset(project_id=project.id, type=AssetType.IMAGE, role=AssetRole.SCENE,
            name="修改版", path="/tmp/revised.png"))
        with patch.object(router.project_service, "revise_scene_reference", new=mocked):
            response = self.client.post(base, json={
                "source_asset_id": "asset_old", "user_suggestions": "改为清晨", "expected_version": 1,
            })
        self.assertEqual(response.status_code, 200, response.text)
        mocked.assert_awaited_once_with(project.id, profile.id, "asset_old", "改为清晨", 1, None, None)

    def test_script_character_backfill_route(self) -> None:
        project = router.project_service.create_project(ProjectBrief(title="角色补足接口", story="甲遇到乙"))
        with patch.object(
            router.project_service,
            "analyze_and_backfill_script_characters",
            new=AsyncMock(return_value={
                "project": project,
                "added_characters": [],
                "added_count": 0,
                "existing_count": 0,
                "message": "当前剧本中的角色均已在角色库，已有角色资料保持不变",
            }),
        ) as analyze:
            response = self.client.post(
                f"/api/projects/{project.id}/characters/analyze-script",
                json={"user_suggestions": "只补充缺少的角色"},
            )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["message"], "当前剧本中的角色均已在角色库，已有角色资料保持不变")
        analyze.assert_awaited_once_with(project.id, "只补充缺少的角色")

    def test_insert_shot_route_forwards_before_first_anchor(self) -> None:
        project = router.project_service.create_project(ProjectBrief(title="第二集转场", story="承接上一集"))
        first = router.store.save_shot(
            Shot(project_id=project.id, ordinal=1, title="原首镜", narrative="第二集正式开场")
        )
        inserted = Shot(project_id=project.id, ordinal=1, title="跨集过渡", narrative="上一集画面淡出")
        mocked = AsyncMock(return_value=inserted)

        with patch.object(router.project_service, "insert_shot_with_ai", new=mocked):
            response = self.client.post(
                f"/api/projects/{project.id}/shots/insert-ai",
                json={
                    "after_shot_id": None,
                    "before_shot_id": first.id,
                    "user_suggestions": "在第一个分镜前增加跨集过渡",
                    "prompt_targets": [],
                },
            )

        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["title"], "跨集过渡")
        mocked.assert_awaited_once_with(
            project.id,
            None,
            "在第一个分镜前增加跨集过渡",
            [],
            "h3-prompt-writing",
            before_shot_id=first.id,
        )

    def test_scene_reference_can_be_uploaded_and_cleared_without_ai(self) -> None:
        project = router.project_service.create_project(ProjectBrief(title="本地场景图", story="人物在客厅"))
        profile = SceneProfile(name="客厅", description="固定窗户和木桌")
        project.scene_profiles = [profile]
        router.store.save_project(project)
        base = f"/api/projects/{project.id}/scene-profiles/{profile.id}"

        uploaded = self.client.post(
            base + "/reference/upload",
            files={"file": ("living-room.png", b"fake-png", "image/png")},
        )
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        payload = uploaded.json()
        self.assertEqual(payload["profile"]["reference_source"], "upload")
        self.assertEqual(payload["profile"]["reference_status"], "completed")
        self.assertTrue(payload["profile"]["approved"])
        self.assertEqual(payload["profile"]["reference_asset_ids"], [payload["asset"]["id"]])
        self.assertEqual(payload["asset"]["role"], "scene")
        self.assertTrue(Path(payload["asset"]["path"]).is_file())
        first_asset_path = payload["asset"]["path"]

        replacement = self.client.post(
            base + "/reference/upload",
            files={"file": ("living-room-v2.webp", b"fake-webp", "image/webp")},
        )
        self.assertEqual(replacement.status_code, 200, replacement.text)
        self.assertEqual(
            replacement.json()["profile"]["reference_asset_ids"],
            [replacement.json()["asset"]["id"]],
        )
        self.assertTrue(Path(first_asset_path).is_file())

        cleared = self.client.delete(base + "/reference")
        self.assertEqual(cleared.status_code, 200, cleared.text)
        self.assertEqual(cleared.json()["reference_asset_ids"], [])
        self.assertEqual(cleared.json()["reference_source"], "none")
        self.assertFalse(cleared.json()["approved"])
        self.assertTrue(Path(first_asset_path).is_file())
        self.assertTrue(Path(replacement.json()["asset"]["path"]).is_file())

    def test_keyframe_prompt_patch_preserves_other_prompt_fields(self) -> None:
        project = self.client.post("/api/projects", json={"title": "首帧局部保存", "story": "人物开门"}).json()
        shot = router.store.save_shot(Shot(
            project_id=project["id"], ordinal=1,
            h3_prompt_skill_output="saved H3", video_prompt="saved H3",
            video_prompt_source="saved H3", seedance_prompt="saved Seedance",
        ))
        url = f"/api/projects/{project['id']}/shots/{shot.id}/keyframe-prompt"
        response = self.client.patch(url, json={
            "keyframe_prompt": "  新首帧 Prompt  ",
            "keyframe_revision_suggestion_draft": "增加逆光",
            "keyframe_revision_mode": "iterate",
        })
        self.assertEqual(response.status_code, 200, response.text)
        saved = self.client.get(f"/api/projects/{project['id']}").json()["shots"][0]
        self.assertEqual(saved["keyframe_prompt"], "新首帧 Prompt")
        self.assertEqual(saved["keyframe_revision_suggestion_draft"], "增加逆光")
        self.assertEqual(saved["keyframe_revision_mode"], "iterate")
        self.assertEqual(saved["h3_prompt_skill_output"], "saved H3")
        self.assertEqual(saved["video_prompt"], "saved H3")
        self.assertEqual(saved["video_prompt_source"], "saved H3")
        self.assertEqual(saved["seedance_prompt"], "saved Seedance")
        self.assertIsNone(saved["keyframe_reference_asset_ids"])
        self.assertGreater(saved["version"], shot.version)
        response = self.client.patch(url, json={
            "keyframe_prompt": "新首帧 Prompt",
            "keyframe_reference_asset_ids": [],
        })
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["keyframe_reference_asset_ids"], [])
        materials = self.client.get(
            f"/api/projects/{project['id']}/shots/{shot.id}/keyframe-materials"
        )
        self.assertEqual(materials.status_code, 200, materials.text)
        self.assertEqual(materials.json()["selection_mode"], "manual")
        self.assertEqual(self.client.patch(url, json={"keyframe_prompt": " "}).status_code, 400)
        self.assertEqual(self.client.patch(url, json={"keyframe_prompt": "x" * 12001}).status_code, 422)
        self.assertEqual(self.client.patch(url, json={"keyframe_prompt": "x", "keyframe_revision_mode": "bad"}).status_code, 422)
        self.assertEqual(self.client.patch(
            f"/api/projects/another-project/shots/{shot.id}/keyframe-prompt",
            json={"keyframe_prompt": "x"},
        ).status_code, 404)

    def test_project_cover_route_forwards_suggestion_and_ratio(self) -> None:
        project = self.client.post(
            "/api/projects",
            json={"title": "排插安全科普", "story": "阴差上门排查老化排插"},
        ).json()
        uploaded = self.client.post(
            f"/api/projects/{project['id']}/assets",
            files={"file": ("layout.png", io.BytesIO(b"cover-reference"), "image/png")},
            data={"role": "cover_reference", "name": "客户封面参考"},
        )
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        self.assertEqual(uploaded.json()["role"], "cover_reference")
        reference_asset_id = uploaded.json()["id"]
        cover = Asset(
            project_id=project["id"],
            type=AssetType.IMAGE,
            role=AssetRole.COVER,
            name="排插安全科普 · 项目封面 9:16",
            path="/tmp/generated-cover.png",
            tags=["generated-cover", "cover-ratio:9:16"],
        )
        generator = AsyncMock(return_value=cover)
        with patch.object(router.project_service, "generate_project_cover", new=generator):
            response = self.client.post(
                f"/api/projects/{project['id']}/cover/generate",
                json={
                    "user_suggestions": "让危险排插占前景三分之一",
                    "aspect_ratio": "9:16",
                    "reference_mode": "uploaded",
                    "reference_asset_id": reference_asset_id,
                },
            )

        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["role"], "cover")
        self.assertIn("cover-ratio:9:16", response.json()["tags"])
        generator.assert_awaited_once_with(
            project_id=project["id"],
            user_suggestions="让危险排插占前景三分之一",
            aspect_ratio="9:16",
            reference_mode="uploaded",
            reference_asset_id=reference_asset_id,
            image_provider=None,
            image_model=None,
        )
        invalid = self.client.post(
            f"/api/projects/{project['id']}/cover/generate",
            json={"aspect_ratio": "2:1"},
        )
        self.assertEqual(invalid.status_code, 422)
        invalid_mode = self.client.post(
            f"/api/projects/{project['id']}/cover/generate",
            json={"reference_mode": "random"},
        )
        self.assertEqual(invalid_mode.status_code, 422)

    def test_shot_split_preview_and_confirm_routes(self) -> None:
        project = self.client.post("/api/projects", json={"title": "拆分接口", "story": "人物起身、开门并走出房间"}).json()
        source = router.store.save_shot(Shot(
            project_id=project["id"],
            ordinal=1,
            title="连续动作",
            narrative="人物起身、开门并走出房间",
            duration_seconds=6,
            video_path="/tmp/source.mp4",
            video_status="completed",
        ))
        tail = router.store.save_shot(Shot(project_id=project["id"], ordinal=2, title="下一镜", narrative="人物来到走廊"))

        class LLM:
            async def generate_json(self, *_args, **_kwargs):
                return {
                    "rationale": "每镜一个主动作",
                    "segments": [
                        {"title": "起身", "duration_seconds": 2, "narrative": "人物起身", "scene_description": "人物坐在椅上"},
                        {"title": "开门走出", "duration_seconds": 4, "narrative": "人物开门走出房间", "scene_description": "人物站在门前"},
                    ],
                }

        base = f"/api/projects/{project['id']}/shots/{source.id}"
        with patch("src.video_workflow.services.projects.create_llm_generator", return_value=LLM()):
            response = self.client.post(base + "/split-preview", json={"segment_count": 2, "user_suggestions": "保持六秒"})
        self.assertEqual(response.status_code, 200, response.text)
        preview = response.json()
        self.assertEqual(len(preview["segments"]), 2)
        self.assertEqual(preview["proposed_duration_seconds"], 6)
        self.assertEqual([shot.id for shot in router.store.list_shots(project["id"])], [source.id, tail.id])

        preview["segments"][0]["title"] = "用户调整后起身"
        running = router.store.save_job(RenderJob(
            project_id=project["id"],
            shot_id=source.id,
            type=JobType.VIDEO,
            status=JobStatus.RUNNING,
        ))
        blocked = self.client.post(base + "/split-confirm", json={"preview": preview, "prompt_targets": []})
        self.assertEqual(blocked.status_code, 409, blocked.text)
        self.assertIn("视频生成任务", blocked.json()["detail"])
        router.store.delete_job(running.id)
        response = self.client.post(base + "/split-confirm", json={"preview": preview, "prompt_targets": []})
        self.assertEqual(response.status_code, 200, response.text)
        created = response.json()
        self.assertEqual(created[0]["title"], "用户调整后起身")
        board = self.client.get(f"/api/projects/{project['id']}").json()["shots"]
        self.assertEqual([shot["ordinal"] for shot in board], [1, 2, 3])
        self.assertEqual(board[-1]["id"], tail.id)
        self.assertEqual(board[-1]["ordinal"], 3)
        self.assertTrue(all(shot["id"] != source.id for shot in board[:2]))
        self.assertTrue(all(shot["video_path"] is None and shot["video_status"] == "pending" for shot in board[:2]))
        self.assertEqual(self.client.post(base + "/split-confirm", json={"preview": preview}).status_code, 404)

    def test_conflicting_keyframe_operations_return_409_without_writing(self) -> None:
        project = self.client.post("/api/projects", json={"title": "首帧防重", "story": "人物开门"}).json()
        shot = router.store.save_shot(Shot(project_id=project["id"], ordinal=1, keyframe_prompt="保留首帧"))
        with patch.object(router.project_service, "ensure_keyframe_idle", side_effect=KeyframeBusyError("本镜首帧正在生成")):
            response = self.client.patch(
                f"/api/projects/{project['id']}/shots/{shot.id}/keyframe-prompt",
                json={"keyframe_prompt": "重复修改"},
            )
            self.assertEqual(response.status_code, 409, response.text)
            response = self.client.post(f"/api/projects/{project['id']}/keyframes/generate", json={"shot_ids": [shot.id]})
            self.assertEqual(response.status_code, 409, response.text)
            response = self.client.post(
                f"/api/projects/{project['id']}/shots/{shot.id}/keyframe/upload",
                files={"file": ("image.png", b"image", "image/png")},
            )
            self.assertEqual(response.status_code, 409, response.text)
        self.assertEqual(router.store.get_shot(shot.id).keyframe_prompt, "保留首帧")
        self.assertEqual(router.store.list_assets(project["id"]), [])

    def test_keyframe_upload_preserves_prompt_saved_during_file_read(self) -> None:
        project = self.client.post("/api/projects", json={"title": "上传并发", "story": "人物开门"}).json()
        shot = router.store.save_shot(Shot(project_id=project["id"], ordinal=1))

        async def read_upload(_file, _size=-1):
            latest = router.store.get_shot(shot.id)
            latest.h3_prompt_skill_output = "new H3 during upload"
            latest.video_prompt = "new H3 during upload"
            router.store.save_shot(latest)
            return b"image"

        with patch("starlette.datastructures.UploadFile.read", new=read_upload):
            response = self.client.post(
                f"/api/projects/{project['id']}/shots/{shot.id}/keyframe/upload",
                files={"file": ("image.png", b"image", "image/png")},
            )
        self.assertEqual(response.status_code, 200, response.text)
        saved = router.store.get_shot(shot.id)
        self.assertEqual(saved.h3_prompt_skill_output, "new H3 during upload")
        self.assertEqual(saved.video_prompt, "new H3 during upload")
        self.assertEqual(saved.image_status, "completed")
        self.assertIsNotNone(saved.keyframe_asset_id)

    def test_deleting_style_asset_keeps_generated_h3_through_refresh_and_keyframes(self) -> None:
        project_id = self.client.post("/api/projects", json={"title": "保留付费 Prompt", "story": "人物开门"}).json()["id"]
        style = self.client.post(f"/api/projects/{project_id}/assets", data={"role": "style", "name": "旧风格图"},
                                 files={"file": ("style.png", b"style", "image/png")}).json()
        project = router.store.get_project(project_id)
        project.style_profile = StyleProfile(name="手绘", reference_asset_ids=[style["id"]], approved=True)
        router.store.save_project(project)
        shot = router.store.save_shot(Shot(project_id=project_id, ordinal=1, narrative="人物开门", keyframe_prompt="门前"))
        output = "integrated_multimodal_description: A person opens the door.\noverall_soundscape: Room tone.\nnon_diegetic_music: None."

        class LLM:
            async def generate_json(self, *args, **kwargs):
                return {"video_prompt": output}

        class ImageGenerator:
            async def generate_image(self, scene, output_dir, *args, **kwargs):
                path = Path(output_dir) / "image.png"
                path.write_bytes(b"image")
                return str(path)

        with patch("src.video_workflow.services.projects.create_llm_generator", return_value=LLM()):
            generated = self.client.post(
                f"/api/projects/{project_id}/h3-prompts/generate",
                json={"shot_ids": [shot.id], "input_mode": "frame"},
            )
        self.assertEqual(generated.status_code, 200, generated.text)
        deleted = self.client.delete(f"/api/projects/{project_id}/assets/{style['id']}")
        self.assertEqual(deleted.status_code, 200, deleted.text)
        with patch("src.video_workflow.services.projects.create_image_generator", return_value=ImageGenerator()):
            image = self.client.post(f"/api/projects/{project_id}/keyframes/generate", json={"shot_ids": [shot.id]})
        self.assertEqual(image.status_code, 200, image.text)
        for _ in range(3):
            bundle = self.client.get(f"/api/projects/{project_id}").json()
            saved = bundle["shots"][0]
            self.assertEqual(saved["h3_prompt_skill_output"], output)
            self.assertEqual(saved["video_prompt"], output)
            self.assertEqual(saved["h3_director_version"], H3_DIRECTOR_VERSION)
            self.assertLess(saved["h3_prompt_source_revision"], saved["content_revision"])
            self.assertEqual(saved["image_status"], "completed")
            self.assertTrue(saved["keyframe_asset_id"])
            self.assertNotIn("旧风格图", saved["seedance_prompt"])
            self.assertNotIn(style["id"], [asset["id"] for asset in bundle["assets"]])
        planned = router.project_service.plan_shot(project_id, shot.id)
        self.assertEqual(planned.video_prompt, output)
        history = self.client.get(f"/api/projects/{project_id}/shots/{shot.id}/h3-prompt-history")
        self.assertEqual(history.status_code, 200)
        self.assertEqual(len(history.json()), 1)
        self.assertEqual(history.json()[0]["prompt"], output)

    def test_stale_shot_form_returns_conflict_instead_of_overwriting_prompt(self) -> None:
        project_id = self.client.post("/api/projects", json={"title": "旧页面防覆盖", "story": "人物开门"}).json()["id"]
        shot = router.store.save_shot(Shot(project_id=project_id, ordinal=1, video_prompt="旧模板"))
        stale = shot.model_dump(mode="json")
        shot.video_prompt = shot.h3_prompt_skill_output = "newly generated English prompt"
        shot.version += 1
        router.store.save_shot(shot)
        response = self.client.put(f"/api/projects/{project_id}/shots/{shot.id}", json=stale)
        self.assertEqual(response.status_code, 409, response.text)
        self.assertEqual(router.store.get_shot(shot.id).video_prompt, "newly generated English prompt")

    def test_extract_storyboard_rows_from_xlsx(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["镜头标题", "剧情内容", "时长(秒)", "Seedance Prompt"])
        sheet.append(["开场", "人物推门进入", 6, "固定中景，人物缓慢进入"])
        payload = io.BytesIO()
        workbook.save(payload)

        rows = router._extract_storyboard_rows("客户分镜.xlsx", payload.getvalue())

        self.assertEqual(
            rows,
            [
                {
                    "title": "开场",
                    "narrative": "人物推门进入",
                    "duration_seconds": "6",
                    "seedance_prompt": "固定中景，人物缓慢进入",
                }
            ],
        )

    def test_project_shot_asset_and_public_review_flow(self) -> None:
        response = self.client.post(
            "/api/projects",
            json={"title": "API 项目", "story": "测试完整项目接口", "target_duration_seconds": 10},
        )
        self.assertEqual(response.status_code, 200, response.text)
        project = response.json()
        project_id = project["id"]

        shot = Shot(project_id=project_id, ordinal=1, narrative="镜头叙事", duration_seconds=5)
        response = self.client.post(f"/api/projects/{project_id}/shots", json=shot.model_dump(mode="json"))
        self.assertEqual(response.status_code, 200, response.text)
        shot_payload = response.json()
        shot_payload["dialogue"] = "你好"
        shot_payload.update(
            h3_width=1344,
            h3_height=768,
            h3_steps=6,
            h3_scheduler="karras",
            h3_denoise=0.9,
            h3_seed=20260822,
            h3_resolution="2K",
            h3_ratio="9:16",
            h3_context_ir_enabled=True,
        )
        response = self.client.put(f"/api/projects/{project_id}/shots/{shot.id}", json=shot_payload)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["dialogue"], "你好")
        self.assertEqual(response.json()["h3_width"], 1344)
        self.assertEqual(response.json()["h3_steps"], 6)
        self.assertEqual(response.json()["h3_scheduler"], "karras")
        self.assertEqual(response.json()["h3_seed"], 20260822)
        self.assertEqual(response.json()["h3_resolution"], "2K")
        self.assertEqual(response.json()["h3_ratio"], "9:16")
        self.assertTrue(response.json()["h3_context_ir_enabled"])

        response = self.client.post(
            f"/api/projects/{project_id}/assets",
            data={"role": "character", "name": "人物参考"},
            files={"file": ("character.png", b"not-a-real-png", "image/png")},
        )
        self.assertEqual(response.status_code, 200, response.text)
        asset = response.json()
        self.assertEqual(asset["role"], "character")

        response = self.client.post(
            f"/api/projects/{project_id}/script/upload",
            files={"file": ("story.txt", "第一幕：角色走进雨中。".encode(), "text/plain")},
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn("角色走进雨中", response.json()["project"]["brief"]["story"])

        response = self.client.get("/api/settings")
        self.assertEqual(response.status_code, 200, response.text)
        groups = {group["id"]: group for group in response.json()["groups"]}
        self.assertIn("h3_api", groups)
        self.assertTrue(any(field["key"] == "H3_PROVIDER" for field in groups["h3_api"]["fields"]))
        secret_fields = [field for group in response.json()["groups"] for field in group["fields"] if field["secret"]]
        self.assertTrue(secret_fields)
        self.assertTrue(all(field["value"] == "" for field in secret_fields))

        shot_payload = self.client.get(f"/api/projects/{project_id}").json()["shots"][0]
        shot_payload["reference_asset_ids"] = [asset["id"]]
        self.assertEqual(
            self.client.put(f"/api/projects/{project_id}/shots/{shot.id}", json=shot_payload).status_code,
            200,
        )

        response = self.client.get(f"/api/review/{project['review_token']}")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["assets"][0]["path"], "")
        response = self.client.post(
            f"/api/review/{project['review_token']}",
            json={
                "target_type": "storyboard",
                "target_id": project_id,
                "decision": "approved",
                "comment": "确认",
                "reviewer": "客户",
            },
        )
        self.assertEqual(response.status_code, 200, response.text)

        approved_bundle = self.client.get(f"/api/projects/{project_id}").json()
        approved_version = approved_bundle["project"]["storyboard_version"]
        self.assertEqual(approved_bundle["project"]["status"], "storyboard_approved")
        self.assertEqual(approved_bundle["shots"][0]["approval_status"], "approved")

        edited = approved_bundle["shots"][0]
        edited["narrative"] = "客户确认后的新修改"
        response = self.client.put(f"/api/projects/{project_id}/shots/{shot.id}", json=edited)
        self.assertEqual(response.status_code, 200, response.text)
        invalidated = self.client.get(f"/api/projects/{project_id}").json()
        self.assertEqual(invalidated["project"]["status"], "storyboard_draft")
        self.assertGreater(invalidated["project"]["storyboard_version"], approved_version)
        self.assertEqual(invalidated["shots"][0]["approval_status"], "draft")

        response = self.client.delete(f"/api/projects/{project_id}/assets/{asset['id']}")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(self.client.get(f"/api/projects/{project_id}").json()["shots"][0]["reference_asset_ids"], [])

        blank = self.client.post(f"/api/projects/{project_id}/shots/blank", json={})
        self.assertEqual(blank.status_code, 200, blank.text)
        shots = self.client.get(f"/api/projects/{project_id}").json()["shots"]
        response = self.client.post(
            f"/api/projects/{project_id}/storyboard/reorder",
            json={"shot_ids": [shots[1]["id"], shots[0]["id"]]},
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()[0]["id"], shots[1]["id"])

        bundle = self.client.get(f"/api/projects/{project_id}").json()
        self.assertEqual(len(bundle["shots"]), 2)
        self.assertEqual(len(bundle["assets"]), 1)
        self.assertEqual(bundle["assets"][0]["type"], "document")
        self.assertEqual(len(bundle["reviews"]), 1)

        delivery = router.store.save_delivery(
            Delivery(project_id=project_id, output_path=str(Path(self.temp.name) / "final.mp4"))
        )
        response = self.client.post(
            f"/api/review/{project['review_token']}",
            json={
                "target_type": "delivery",
                "target_id": delivery.id,
                "decision": "approved",
                "comment": "成片确认",
                "reviewer": "客户",
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        delivered = self.client.get(f"/api/projects/{project_id}").json()
        self.assertEqual(delivered["project"]["status"], "delivered")
        self.assertEqual(len(delivered["reviews"]), 2)

    def test_keyframe_export_only_includes_generated_selected_shots(self) -> None:
        project = self.client.post(
            "/api/projects",
            json={"title": "分镜图导出测试", "story": "两个镜头", "target_duration_seconds": 10},
        ).json()
        project_id = project["id"]
        first = router.store.save_shot(Shot(project_id=project_id, ordinal=1, title="开场/入镜"))
        second = router.store.save_shot(Shot(project_id=project_id, ordinal=2, title="尚未生成"))
        image_path = router.project_service.project_dir(project_id) / "assets" / "first-frame.png"
        image_path.parent.mkdir(parents=True, exist_ok=True)
        image_path.write_bytes(b"fake-png-content")
        asset = router.project_service.register_existing_asset(
            project_id,
            image_path,
            AssetRole.KEYFRAME,
            "镜头 1 首帧",
        )
        first.keyframe_asset_id = asset.id
        first.image_status = "completed"
        router.store.save_shot(first)

        response = self.client.post(
            f"/api/projects/{project_id}/keyframes/export",
            json={"shot_ids": [second.id, first.id]},
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.headers["content-type"], "application/zip")
        self.assertIn("filename*=utf-8", response.headers["content-disposition"].lower())
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            self.assertEqual(archive.namelist(), ["镜头_001_开场_入镜.png"])
            self.assertEqual(archive.read(archive.namelist()[0]), b"fake-png-content")

        empty = self.client.post(
            f"/api/projects/{project_id}/keyframes/export",
            json={"shot_ids": [second.id]},
        )
        self.assertEqual(empty.status_code, 400, empty.text)
        self.assertIn("没有可导出", empty.json()["detail"])

    def test_storyboard_csv_exports_structured_voice_events_with_clear_speakers(self) -> None:
        project = router.project_service.create_project(ProjectBrief(title="对白导出", story="系统播报后角色回应"))
        router.store.save_shot(Shot(
            project_id=project.id,
            ordinal=1,
            narrative="系统宣布规则，角色回应",
            duration_seconds=8,
            voice_events=[
                VoiceEvent(
                    kind="system_vo",
                    speaker_name="系统",
                    text="第一关开始。",
                    start_seconds=0.5,
                    end_seconds=2,
                ),
                VoiceEvent(
                    kind="character",
                    speaker_name="陆峥",
                    text="立即停止作业。",
                    start_seconds=2.2,
                    end_seconds=4,
                    lip_sync=True,
                ),
                VoiceEvent(
                    kind="narration",
                    speaker_name="陈野",
                    text="众人终于意识到风险。",
                    start_seconds=4.2,
                    end_seconds=7,
                ),
            ],
        ))
        router.store.save_shot(Shot(
            project_id=project.id,
            ordinal=2,
            narrative="旧分镜旁白",
            dialogue="危险已经解除。",
            duration_seconds=4,
        ))

        response = self.client.get(f"/api/projects/{project.id}/storyboard.csv")
        self.assertEqual(response.status_code, 200, response.text)
        rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
        self.assertEqual(rows[0][3], "对白")
        self.assertEqual(
            rows[1][3],
            "系统播报：第一关开始。\n陆峥：立即停止作业。\n陈野（旁白）：众人终于意识到风险。",
        )
        self.assertEqual(rows[2][3], "旁白（未指定角色）：危险已经解除。")

        xlsx_response = self.client.get(f"/api/projects/{project.id}/storyboard.xlsx")
        self.assertEqual(xlsx_response.status_code, 200, xlsx_response.text)
        self.assertEqual(
            xlsx_response.headers["content-type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(io.BytesIO(xlsx_response.content), data_only=True)
        sheet = workbook["分镜表"]
        self.assertEqual(sheet["D1"].value, "对白")
        self.assertEqual(
            sheet["D2"].value,
            "系统播报：第一关开始。\n陆峥：立即停止作业。\n陈野（旁白）：众人终于意识到风险。",
        )
        self.assertEqual(sheet["D3"].value, "旁白（未指定角色）：危险已经解除。")

    def test_video_export_downloads_selected_completed_jobs_as_zip(self) -> None:
        project = router.project_service.create_project(ProjectBrief(title="视频批量导出", story="两个已完成视频"))
        first_shot = router.store.save_shot(Shot(project_id=project.id, ordinal=1, title="开场/入镜"))
        second_shot = router.store.save_shot(Shot(project_id=project.id, ordinal=2, title="镜头二"))
        first_path = Path(self.temp.name) / "first.mp4"
        second_path = Path(self.temp.name) / "second.mp4"
        first_path.write_bytes(b"first-video")
        second_path.write_bytes(b"second-video")
        first_job = router.store.save_job(RenderJob(
            project_id=project.id,
            shot_id=first_shot.id,
            type=JobType.VIDEO,
            status=JobStatus.COMPLETED,
            output_path=str(first_path),
        ))
        second_job = router.store.save_job(RenderJob(
            project_id=project.id,
            shot_id=second_shot.id,
            type=JobType.VIDEO,
            status=JobStatus.COMPLETED,
            output_path=str(second_path),
        ))
        ignored = router.store.save_job(RenderJob(
            project_id=project.id,
            shot_id=second_shot.id,
            type=JobType.VIDEO,
            status=JobStatus.FAILED,
            output_path=str(second_path),
        ))

        response = self.client.post(
            f"/api/projects/{project.id}/videos/export",
            json={"job_ids": [second_job.id, ignored.id, first_job.id]},
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.headers["content-type"], "application/zip")
        self.assertIn("filename*=utf-8", response.headers["content-disposition"].lower())
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            names = archive.namelist()
            self.assertEqual(len(names), 2)
            self.assertEqual({archive.read(name) for name in names}, {b"first-video", b"second-video"})
            self.assertTrue(any(first_job.id[-6:] in name for name in names))
            self.assertTrue(any(second_job.id[-6:] in name for name in names))

        empty = self.client.post(
            f"/api/projects/{project.id}/videos/export",
            json={"job_ids": [ignored.id]},
        )
        self.assertEqual(empty.status_code, 400, empty.text)
        self.assertIn("已完成视频", empty.json()["detail"])


    def test_render_job_delete_requires_terminal_status(self) -> None:
        response = self.client.post(
            "/api/projects",
            json={"title": "删除任务", "story": "删除渲染任务记录", "target_duration_seconds": 10},
        )
        self.assertEqual(response.status_code, 200, response.text)
        project_id = response.json()["id"]

        shot = Shot(project_id=project_id, ordinal=1, narrative="镜头叙事", duration_seconds=5)
        router.store.save_shot(shot)
        completed_job = router.store.save_job(
            RenderJob(project_id=project_id, shot_id=shot.id, type=JobType.VIDEO, status=JobStatus.COMPLETED)
        )
        running_job = router.store.save_job(
            RenderJob(project_id=project_id, shot_id=shot.id, type=JobType.VIDEO, status=JobStatus.RUNNING)
        )

        response = self.client.delete(f"/api/projects/{project_id}/jobs/{running_job.id}")
        self.assertEqual(response.status_code, 409)
        self.assertIn("进行中", response.json()["detail"])

        response = self.client.delete(f"/api/projects/{project_id}/jobs/{completed_job.id}")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertTrue(response.json()["deleted"])
        self.assertIsNone(router.store.get_job(completed_job.id))
        # running job remains -> video_status follows the latest remaining job
        self.assertEqual(router.store.get_shot(shot.id).video_status, "running")

        running_job.status = JobStatus.FAILED
        router.store.save_job(running_job)
        response = self.client.delete(f"/api/projects/{project_id}/jobs/{running_job.id}")
        self.assertEqual(response.status_code, 200, response.text)
        # no job left for the shot -> back to pending
        self.assertEqual(router.store.get_shot(shot.id).video_status, "pending")

        response = self.client.delete(f"/api/projects/{project_id}/jobs/{running_job.id}")
        self.assertEqual(response.status_code, 404)

    def test_completed_render_selection_locks_video_and_continuity_tail(self) -> None:
        project = router.project_service.create_project(ProjectBrief(title="抽卡选片", story="连续两镜"))
        video_a = Path(self.temp.name) / "card-a.mp4"
        video_b = Path(self.temp.name) / "card-b.mp4"
        tail_a_path = Path(self.temp.name) / "card-a-tail.png"
        tail_b_path = Path(self.temp.name) / "card-b-tail.png"
        video_a.write_bytes(b"video-a")
        video_b.write_bytes(b"video-b")
        tail_a_path.write_bytes(b"tail-a")
        tail_b_path.write_bytes(b"tail-b")
        tail_a = router.project_service.register_existing_asset(project.id, tail_a_path, AssetRole.LAST_FRAME, "A 尾帧")
        tail_b = router.project_service.register_existing_asset(project.id, tail_b_path, AssetRole.LAST_FRAME, "B 尾帧")
        first = router.store.save_shot(Shot(project_id=project.id, ordinal=1, title="第一镜"))
        second = router.store.save_shot(Shot(
            project_id=project.id,
            ordinal=2,
            title="第二镜",
            continuity_mode=ShotContinuityMode.CONTINUOUS,
            continuity_source_shot_id=first.id,
        ))
        card_a = router.store.save_job(RenderJob(
            project_id=project.id,
            shot_id=first.id,
            type=JobType.VIDEO,
            status=JobStatus.COMPLETED,
            provider="ark_seedance",
            output_path=str(video_a),
            input_snapshot={"seedance_result": {"last_frame_asset_id": tail_a.id}},
        ))
        card_b = router.store.save_job(RenderJob(
            project_id=project.id,
            shot_id=first.id,
            type=JobType.VIDEO,
            status=JobStatus.COMPLETED,
            provider="ark_seedance",
            output_path=str(video_b),
            input_snapshot={"output_artifacts": {"last_frame_asset_id": tail_b.id}},
        ))
        first.video_path = card_b.output_path
        first.last_frame_asset_id = tail_b.id
        first.video_status = "completed"
        router.store.save_shot(first)

        response = self.client.post(f"/api/projects/{project.id}/jobs/{card_a.id}/select", json={})
        self.assertEqual(response.status_code, 200, response.text)
        selected = router.store.get_shot(first.id)
        self.assertEqual(selected.selected_video_job_id, card_a.id)
        self.assertEqual(selected.video_path, str(video_a))
        self.assertEqual(selected.last_frame_asset_id, tail_a.id)
        self.assertEqual(
            router.project_service.continuity_input_asset_id(project, second),
            tail_a.id,
        )

        # A newly queued or completed candidate does not replace the adopted card.
        router.render_queue._set_job_status_on_shot(selected, "queued")
        self.assertEqual(selected.video_status, "completed")
        candidate = RenderJob(
            project_id=project.id,
            shot_id=first.id,
            type=JobType.VIDEO,
            status=JobStatus.COMPLETED,
            output_path=str(video_b),
        )
        router.render_queue._bind_completed_output(
            selected,
            candidate,
            last_frame_asset_id=tail_b.id,
            replace_tail=True,
        )
        self.assertEqual(selected.video_path, str(video_a))
        self.assertEqual(selected.last_frame_asset_id, tail_a.id)

        # Render metadata is server-owned and survives an ordinary editor save.
        edited = selected.model_copy(update={"title": "第一镜改名"})
        saved = router.project_service.update_shot(project.id, edited)
        self.assertEqual(saved.selected_video_job_id, card_a.id)
        self.assertEqual(saved.last_frame_asset_id, tail_a.id)

        delete_selected = self.client.delete(f"/api/projects/{project.id}/jobs/{card_a.id}")
        self.assertEqual(delete_selected.status_code, 409, delete_selected.text)
        self.assertIn("当前采用版本", delete_selected.json()["detail"])

        incomplete = router.store.save_job(RenderJob(
            project_id=project.id,
            shot_id=first.id,
            type=JobType.VIDEO,
            status=JobStatus.FAILED,
        ))
        rejected = self.client.post(f"/api/projects/{project.id}/jobs/{incomplete.id}/select", json={})
        self.assertEqual(rejected.status_code, 400, rejected.text)

    def test_keyframe_upload_replaces_shot_binding(self) -> None:
        project = self.client.post(
            "/api/projects",
            json={"title": "首帧上传测试", "story": "一个镜头", "target_duration_seconds": 5},
        ).json()
        project_id = project["id"]
        shot = router.store.save_shot(Shot(project_id=project_id, ordinal=1, title="开场"))

        response = self.client.post(
            f"/api/projects/{project_id}/shots/{shot.id}/keyframe/upload",
            files={"file": ("local-frame.png", b"fake-png-content", "image/png")},
        )
        self.assertEqual(response.status_code, 200, response.text)
        updated = response.json()
        self.assertEqual(updated["image_status"], "completed")
        self.assertTrue(updated["keyframe_asset_id"])
        asset = router.store.get_asset(updated["keyframe_asset_id"])
        self.assertIsNotNone(asset)
        self.assertEqual(asset.role, AssetRole.KEYFRAME)
        self.assertIn("本地上传", asset.name)
        reloaded = router.store.get_shot(shot.id)
        self.assertEqual(reloaded.keyframe_asset_id, asset.id)
        self.assertTrue(Path(reloaded.image_path).exists())

        bad = self.client.post(
            f"/api/projects/{project_id}/shots/{shot.id}/keyframe/upload",
            files={"file": ("notes.txt", b"not an image", "text/plain")},
        )
        self.assertEqual(bad.status_code, 400)
        self.assertIn("仅支持", bad.json()["detail"])


if __name__ == "__main__":
    unittest.main()
