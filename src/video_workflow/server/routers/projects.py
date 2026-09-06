from __future__ import annotations

import asyncio
import csv
import io
import logging
import mimetypes
import re
import shutil
import tempfile
import zipfile
from pathlib import Path
from typing import Any, Literal
from uuid import uuid4

from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel, Field
from starlette.background import BackgroundTask

from src.video_workflow.config import settings
from src.video_workflow.domain import (
    ApprovalStatus,
    AssetRole,
    AssetType,
    CharacterProfile,
    FinalizeRequest,
    GenerationMode,
    JobStatus,
    Project,
    ProjectBrief,
    ProjectStatus,
    Review,
    Shot,
    ShotSplitPreview,
)
from src.video_workflow.integrations.comfyui import (
    ComfyUIClient,
    H3WorkflowBuilder,
    h3_diffusion_model,
    h3_text_encoder,
    resolve_h3_model_profile,
    resolve_h3_text_encoder_profile,
)
from src.video_workflow.integrations.metaso_h3 import MetaSoH3Client
from src.video_workflow.integrations.seedance import (
    SeedanceClient,
    estimate_seedance_cost,
    resolve_seedance_model,
    seedance_catalog,
    validate_seedance_resolution,
    verify_seedance_asset_signature,
)
from src.video_workflow.h3_prompt_skills import list_h3_prompt_skills
from src.video_workflow.media_paths import resolve_media_path
from src.video_workflow.services.finalize import Finalizer
from src.video_workflow.services.projects import H3_DIRECTOR_VERSION, KeyframeBusyError, ProjectService, SceneReferenceConflictError, ShotSplitBusyError, ShotVersionConflictError
from src.video_workflow.services.render_queue import RenderQueue
from src.video_workflow.storage import ProjectStore

router = APIRouter(prefix="/projects", tags=["projects"])
logger = logging.getLogger(__name__)
public_router = APIRouter(prefix="/review", tags=["public-review"])

store = ProjectStore(settings.DATABASE_PATH)
project_service = ProjectService(store)
render_queue = RenderQueue(store, project_service)
finalizer = Finalizer(store, project_service)


class StoryboardGenerateRequest(BaseModel):
    shot_count: int | None = Field(default=None, ge=1, le=500)
    count_mode: Literal["manual", "ai"] = "manual"
    user_suggestions: str = Field(default="", max_length=4000)
    prompt_targets: list[Literal["h3", "seedance"]] | None = None
    h3_skill_id: str = Field(default="h3-prompt-writing", min_length=1, max_length=100)


class StyleAnalyzeRequest(BaseModel):
    asset_ids: list[str] | None = None
    apply: bool = False


class SceneProfilesGenerateRequest(BaseModel):
    user_suggestions: str = Field(default="", max_length=4000)


class SceneProfilesApplyRequest(BaseModel):
    scene_profile_ids: list[str] | None = None


class SceneReferenceGenerateRequest(BaseModel):
    user_suggestions: str = Field(default="", max_length=4000)
    image_provider: str | None = None
    image_model: str | None = None
    prompt: str | None = Field(default=None, max_length=12000)
    expected_version: int | None = Field(default=None, ge=1)


class SceneProfileUpdateRequest(BaseModel):
    expected_version: int = Field(ge=1)
    name: str = Field(min_length=1, max_length=200)
    description: str = Field(min_length=1, max_length=12000)
    continuity_notes: str = Field(default="", max_length=12000)
    reference_prompt: str = Field(default="", max_length=12000)


class CharacterReferencesGenerateRequest(BaseModel):
    character_ids: list[str] | None = None
    user_suggestions: str = Field(default="", max_length=4000)
    image_provider: str | None = None
    image_model: str | None = None
    reference_asset_ids: list[str] | None = None
    appearance_profile_id: str | None = None
    reference_strategy: Literal["identity", "project_style"] = "identity"


class MissingCharacterReferencesGenerateRequest(BaseModel):
    mode: Literal["script_style", "complete_missing"] = "script_style"
    user_suggestions: str = Field(default="", max_length=4000)
    image_provider: str | None = None
    image_model: str | None = None
    reference_asset_ids: list[str] | None = None


class CharacterReferenceAnalyzeRequest(BaseModel):
    character_id: str
    reference_asset_ids: list[str] = Field(min_length=1, max_length=10)
    appearance_profile_id: str | None = None
    appearance_label: str = Field(default="", max_length=100)
    user_suggestions: str = Field(default="", max_length=4000)


class ShotReviseRequest(BaseModel):
    user_suggestions: str = Field(min_length=1, max_length=4000)
    prompt_targets: list[Literal["h3", "seedance"]] = Field(default_factory=list)
    h3_skill_id: str = Field(default="h3-prompt-writing", min_length=1, max_length=100)


class ShotInsertRequest(BaseModel):
    after_shot_id: str | None = None
    user_suggestions: str = Field(min_length=1, max_length=4000)
    prompt_targets: list[Literal["h3", "seedance"]] = Field(default_factory=lambda: ["seedance"])
    h3_skill_id: str = Field(default="h3-prompt-writing", min_length=1, max_length=100)


class ShotSplitPreviewRequest(BaseModel):
    user_suggestions: str = Field(default="", max_length=4000)
    segment_count: Literal[2, 3, 4] | None = None


class ShotSplitConfirmRequest(BaseModel):
    preview: ShotSplitPreview
    prompt_targets: list[Literal["h3", "seedance"]] = Field(default_factory=list)
    h3_skill_id: str = Field(default="h3-prompt-writing", min_length=1, max_length=100)


class ScriptRewriteRequest(BaseModel):
    mode: Literal["auto", "expand", "shorten"] = "auto"
    user_suggestions: str = Field(default="", max_length=4000)


class SeriesSaveRequest(BaseModel):
    name: str = Field(default="", max_length=200)
    description: str = Field(default="", max_length=2000)


class SeriesEpisodeCreateRequest(BaseModel):
    brief: ProjectBrief
    character_ids: list[str] | None = None


class StoryboardApprovalRequest(BaseModel):
    approved: bool
    comment: str = ""
    reviewer: str = ""


class RenderRequest(BaseModel):
    shot_ids: list[str] | None = None
    provider: Literal["comfyui_h3", "metaso_h3", "atlas_h3", "ark_seedance"] = "comfyui_h3"
    model_id: str | None = None
    resolution: str | None = None
    generate_audio: bool = True


class KeyframeRequest(BaseModel):
    shot_ids: list[str] | None = None
    image_provider: str | None = None
    image_model: str | None = None
    revision_mode: Literal["fresh", "iterate"] = "fresh"
    user_suggestions: str = Field(default="", max_length=4000)


class KeyframeExportRequest(BaseModel):
    shot_ids: list[str] | None = None


class KeyframePromptUpdateRequest(BaseModel):
    keyframe_prompt: str = Field(min_length=1, max_length=12000)
    keyframe_revision_suggestion_draft: str | None = Field(default=None, max_length=4000)
    keyframe_revision_mode: Literal["fresh", "iterate"] | None = None
    keyframe_reference_asset_ids: list[str] | None = Field(default=None, max_length=20)


class H3PromptGenerateRequest(BaseModel):
    shot_ids: list[str] | None = None
    skill_id: str = Field(default="h3-prompt-writing", min_length=1, max_length=100)
    user_suggestions: str = Field(default="", max_length=4000)


class SeedancePromptGenerateRequest(BaseModel):
    shot_ids: list[str] | None = None


class SeedanceEstimateRequest(BaseModel):
    shot_ids: list[str] | None = None
    model_id: str
    resolution: str


class ReorderRequest(BaseModel):
    shot_ids: list[str]


class PublicReviewRequest(BaseModel):
    target_type: str = "storyboard"
    target_id: str
    decision: ApprovalStatus
    comment: str = ""
    reviewer: str = ""


def _not_found(message: str) -> HTTPException:
    return HTTPException(status_code=404, detail=message)


def _bundle(project_id: str) -> dict[str, Any]:
    project = store.get_project(project_id)
    if project is None:
        raise _not_found("Project not found")
    project = project_service.reconcile_scene_reference_status(project)
    return {
        "project": project,
        # Older storyboards matched only the full character-card name, so
        # bracketed aliases such as “女试炼者A（…）” were missing from the shot
        # checkboxes even though the prompt visibly contained that role.
        "shots": project_service.synchronize_storyboard_casts(project_id),
        "assets": store.list_assets(project_id),
        "jobs": store.list_jobs(project_id),
        "reviews": store.list_reviews(project_id),
        "deliveries": store.list_deliveries(project_id),
    }


@router.get("")
async def list_projects():
    return store.list_projects()


@router.post("")
async def create_project(brief: ProjectBrief):
    return project_service.create_project(brief)


@router.get("/series")
async def list_production_series():
    return store.list_series()


@router.post("/series/{series_id}/episodes")
async def create_series_episode(series_id: str, request: SeriesEpisodeCreateRequest):
    try:
        return project_service.create_series_episode(
            series_id,
            request.brief,
            request.character_ids,
        )
    except KeyError as exc:
        raise _not_found(str(exc)) from exc
    except (ValueError, OSError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.get("/h3-prompt-skills")
async def get_h3_prompt_skills():
    return {
        "default_skill_id": "h3-prompt-writing",
        "director_version": H3_DIRECTOR_VERSION,
        "skills": list_h3_prompt_skills(),
    }


@router.get("/seedance/catalog")
async def get_seedance_catalog():
    return seedance_catalog()


@router.get("/seedance-assets/{asset_id}")
async def download_seedance_asset(asset_id: str, expires: int, signature: str):
    """Short-lived, signed media endpoint consumed by Volcengine Ark."""
    if not verify_seedance_asset_signature(asset_id, expires, signature):
        raise HTTPException(status_code=403, detail="Seedance asset link expired or invalid")
    asset = store.get_asset(asset_id)
    path = resolve_media_path(asset.path) if asset else None
    if path is None or not path.is_file():
        raise _not_found("Seedance asset not found")
    guessed = mimetypes.guess_type(path.name)[0]
    media_type = asset.mime_type if asset.mime_type != "application/octet-stream" else guessed
    return FileResponse(
        path,
        media_type=media_type or "application/octet-stream",
        headers={"Cache-Control": "private, max-age=300"},
    )


@router.get("/legacy/sessions")
async def list_legacy_sessions():
    return project_service.list_legacy_sessions()


@router.post("/legacy/import/{session_id}")
async def import_legacy_session(session_id: str):
    try:
        return project_service.import_legacy_session(session_id)
    except FileNotFoundError as exc:
        raise _not_found(str(exc)) from exc
    except (ValueError, OSError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.get("/{project_id}")
async def get_project(project_id: str):
    return _bundle(project_id)


@router.put("/{project_id}")
async def update_project(project_id: str, project: Project):
    if project.id != project_id:
        raise HTTPException(status_code=400, detail="Project id mismatch")
    if store.get_project(project_id) is None:
        raise _not_found("Project not found")
    return project_service.update_project(project)


@router.post("/{project_id}/series")
async def save_project_as_series(project_id: str, request: SeriesSaveRequest):
    try:
        series, project = project_service.save_project_as_series(
            project_id,
            request.name,
            request.description,
        )
        return {"series": series, "project": project}
    except KeyError as exc:
        raise _not_found(str(exc)) from exc
    except (ValueError, OSError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.delete("/{project_id}")
async def delete_project(project_id: str):
    project_dir = project_service.project_dir(project_id).resolve()
    projects_root = settings.PROJECTS_DIR.resolve()
    if not store.delete_project(project_id):
        raise _not_found("Project not found")
    if project_dir.exists() and projects_root in project_dir.parents:
        shutil.rmtree(project_dir)
    return {"deleted": True}


@router.post("/{project_id}/storyboard/generate")
async def generate_storyboard(project_id: str, request: StoryboardGenerateRequest):
    try:
        shots = await project_service.generate_storyboard(
            project_id,
            request.shot_count,
            request.count_mode,
            request.user_suggestions,
            request.prompt_targets,
            request.h3_skill_id,
        )
        return {"shots": shots, "project": store.get_project(project_id)}
    except KeyError as exc:
        raise _not_found(str(exc)) from exc
    except ValueError as exc:
        logger.warning("项目 %s 生成分镜未通过校验: %s", project_id, exc)
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        logger.exception("项目 %s 生成分镜失败", project_id)
        raise HTTPException(status_code=500, detail=str(exc)) from exc


@router.post("/{project_id}/brief/analyze")
async def analyze_brief(project_id: str):
    try:
        return await project_service.analyze_brief(project_id)
    except KeyError as exc:
        raise _not_found(str(exc)) from exc
    except ValueError as exc:
        logger.warning("项目 %s 的 AI 剧本分析未通过校验: %s", project_id, exc)
        raise HTTPException(status_code=400, detail=f"AI 剧本分析失败: {exc}") from exc
    except Exception as exc:
        logger.exception("项目 %s 的 AI 剧本分析失败", project_id)
        raise HTTPException(status_code=500, detail=f"AI 剧本分析失败: {exc}") from exc


@router.post("/{project_id}/style/analyze")
async def analyze_style(project_id: str, request: StyleAnalyzeRequest):
    try:
        return await project_service.analyze_style(project_id, request.asset_ids, apply=request.apply)
    except KeyError as exc:
        raise _not_found(str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        logger.exception("项目 %s 的参考风格分析失败", project_id)
        raise HTTPException(status_code=500, detail=f"参考风格分析失败: {exc}") from exc


@router.post("/{project_id}/scene-profiles/generate")
async def generate_scene_profiles(project_id: str, request: SceneProfilesGenerateRequest):
    try:
        return await project_service.generate_scene_profiles(project_id, request.user_suggestions)
    except (KeyError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.post("/{project_id}/scene-profiles/apply")
async def apply_scene_profiles(project_id: str, request: SceneProfilesApplyRequest):
    try:
        return project_service.apply_scene_profiles_to_shots(project_id, request.scene_profile_ids)
    except KeyError as exc:
        raise _not_found(str(exc)) from exc
    except KeyframeBusyError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.get("/{project_id}/scene-profiles/{scene_profile_id}/reference-prompt")
async def get_scene_reference_prompt(project_id: str, scene_profile_id: str):
    try:
        return project_service.scene_reference_prompt(project_id, scene_profile_id)
    except KeyError as exc:
        raise _not_found(str(exc)) from exc


@router.patch("/{project_id}/scene-profiles/{scene_profile_id}")
async def update_scene_profile(project_id: str, scene_profile_id: str, request: SceneProfileUpdateRequest):
    try:
        return project_service.update_scene_profile(project_id, scene_profile_id,
            request.model_dump(exclude={"expected_version"}), request.expected_version)
    except KeyError as exc:
        raise _not_found(str(exc)) from exc
    except SceneReferenceConflictError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.post("/{project_id}/scene-profiles/{scene_profile_id}/reference-prompt")
async def preview_scene_reference_prompt(project_id: str, scene_profile_id: str, request: SceneProfileUpdateRequest):
    try:
        project, profile = project_service.require_scene_profile(project_id, scene_profile_id)
        draft = profile.model_copy(update=request.model_dump(exclude={"expected_version"}))
        return {"prompt": project_service.compile_scene_reference_prompt(project, draft)}
    except KeyError as exc:
        raise _not_found(str(exc)) from exc


@router.post("/{project_id}/scene-profiles/{scene_profile_id}/reference/retry-download")
async def retry_scene_reference_download(project_id: str, scene_profile_id: str):
    try:
        return await project_service.retry_scene_reference_download(project_id, scene_profile_id)
    except KeyError as exc:
        raise _not_found(str(exc)) from exc
    except SceneReferenceConflictError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        logger.exception("场景母版图下载恢复失败: %s", scene_profile_id)
        raise HTTPException(status_code=502, detail=str(exc) or type(exc).__name__) from exc


@router.post("/{project_id}/scene-profiles/{scene_profile_id}/reference")
async def generate_scene_reference(
    project_id: str,
    scene_profile_id: str,
    request: SceneReferenceGenerateRequest,
):
    try:
        return await project_service.generate_scene_reference(
            project_id,
            scene_profile_id,
            request.user_suggestions,
            request.image_provider,
            request.image_model,
            prompt=request.prompt,
            expected_version=request.expected_version,
        )
    except KeyError as exc:
        raise _not_found(str(exc)) from exc
    except SceneReferenceConflictError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        logger.exception("场景母版图生成或下载失败: %s", scene_profile_id)
        raise HTTPException(status_code=502, detail=str(exc) or type(exc).__name__) from exc


@router.post("/{project_id}/characters/references/analyze")
async def analyze_character_references(project_id: str, request: CharacterReferenceAnalyzeRequest):
    try:
        return await project_service.analyze_character_references(
            project_id,
            request.character_id,
            request.reference_asset_ids,
            request.appearance_profile_id,
            request.appearance_label,
            request.user_suggestions,
        )
    except KeyError as exc:
        raise _not_found(str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.post("/{project_id}/characters/references/generate")
async def generate_character_references(project_id: str, request: CharacterReferencesGenerateRequest):
    try:
        return await project_service.generate_character_references(
            project_id,
            request.character_ids,
            request.user_suggestions,
            request.image_provider,
            request.image_model,
            request.reference_asset_ids,
            request.appearance_profile_id,
            reference_strategy=request.reference_strategy,
        )
    except (KeyError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.post("/{project_id}/characters/references/generate-missing")
async def generate_missing_character_references(
    project_id: str,
    request: MissingCharacterReferencesGenerateRequest,
):
    try:
        return await project_service.generate_missing_character_references(
            project_id,
            mode=request.mode,
            user_suggestions=request.user_suggestions,
            image_provider=request.image_provider,
            image_model=request.image_model,
            reference_asset_ids=request.reference_asset_ids,
        )
    except (KeyError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.post("/{project_id}/brief/rewrite")
async def rewrite_brief(project_id: str, request: ScriptRewriteRequest):
    try:
        return await project_service.rewrite_story(project_id, request.mode, request.user_suggestions)
    except KeyError as exc:
        raise _not_found(str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"AI 剧本改写失败: {exc}") from exc


def _extract_script_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md", ".markdown"}:
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                return content.decode(encoding).strip()
            except UnicodeDecodeError:
                continue
        raise ValueError("文本编码无法识别，请转换为 UTF-8 后重试")
    if suffix == ".docx":
        from docx import Document
        document = Document(io.BytesIO(content))
        return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()).strip()
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        return "\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()
    raise ValueError("支持 TXT、Markdown、DOCX 和可提取文字的 PDF 剧本")


_STORYBOARD_HEADER_ALIASES = {
    "镜头标题": "title", "标题": "title", "title": "title",
    "画面叙事": "narrative", "剧情": "narrative", "剧情内容": "narrative",
    "分镜内容": "narrative", "镜头内容": "narrative", "内容": "narrative", "narrative": "narrative",
    "对白": "dialogue", "台词": "dialogue", "旁白": "dialogue", "dialogue": "dialogue",
    "发言者": "dialogue_speaker", "说话人": "dialogue_speaker", "dialogue_speaker": "dialogue_speaker",
    "场景": "scene_description", "场景描述": "scene_description", "scene_description": "scene_description",
    "角色": "character_names", "出场人物": "character_names", "人物": "character_names", "character_names": "character_names",
    "景别": "shot_size", "shot_size": "shot_size",
    "机位": "camera_angle", "角度": "camera_angle", "机位/角度": "camera_angle", "camera_angle": "camera_angle",
    "镜头": "lens", "焦段": "lens", "镜头/焦段": "lens", "lens": "lens",
    "运镜": "camera_motion", "camera_motion": "camera_motion",
    "主体动作": "subject_motion", "动作": "subject_motion", "subject_motion": "subject_motion",
    "转场": "transition", "transition": "transition",
    "声音设计": "audio_design", "音频": "audio_design", "audio_design": "audio_design",
    "时长": "duration_seconds", "时长(秒)": "duration_seconds", "时长（秒）": "duration_seconds", "duration": "duration_seconds", "duration_seconds": "duration_seconds",
    "生成策略": "generation_mode", "generation_mode": "generation_mode",
    "通用prompt": "visual_prompt", "通用分镜prompt": "visual_prompt", "visual_prompt": "visual_prompt",
    "首帧prompt": "keyframe_prompt", "首帧图prompt": "keyframe_prompt", "keyframe_prompt": "keyframe_prompt",
    "h3 prompt": "h3_prompt", "h3_prompt": "h3_prompt", "minimax h3 prompt": "h3_prompt",
    "seedance prompt": "seedance_prompt", "seedance_prompt": "seedance_prompt",
}


def _normalize_storyboard_header(value: Any) -> str:
    text = re.sub(r"\s+", " ", str(value or "").strip()).lower()
    compact = text.replace(" ", "")
    return _STORYBOARD_HEADER_ALIASES.get(text) or _STORYBOARD_HEADER_ALIASES.get(compact) or compact


def _extract_storyboard_rows(filename: str, content: bytes) -> list[dict[str, str]]:
    suffix = Path(filename).suffix.lower()
    raw_rows: list[list[Any]] = []
    if suffix == ".csv":
        decoded = ""
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                decoded = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if not decoded:
            raise ValueError("CSV 编码无法识别，请另存为 UTF-8 CSV")
        raw_rows = list(csv.reader(io.StringIO(decoded)))
    elif suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        raw_rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    else:
        raise ValueError("分镜表支持 CSV、XLSX 和 XLSM 文件")
    raw_rows = [row for row in raw_rows if any(str(value or "").strip() for value in row)]
    if len(raw_rows) < 2:
        raise ValueError("分镜表至少需要一行表头和一行分镜数据")
    headers = [_normalize_storyboard_header(value) for value in raw_rows[0]]
    rows: list[dict[str, str]] = []
    for raw in raw_rows[1:]:
        row = {
            header: str(raw[index] or "").strip()
            for index, header in enumerate(headers)
            if header and index < len(raw) and str(raw[index] or "").strip()
        }
        if row:
            rows.append(row)
    return rows


@router.post("/{project_id}/storyboard/import")
async def import_storyboard(
    project_id: str,
    file: UploadFile = File(...),
    user_suggestions: str = Form(""),
    prompt_targets: str = Form("seedance"),
    h3_skill_id: str = Form("h3-prompt-writing"),
):
    project_service.require_project(project_id)
    filename = Path(file.filename or "storyboard.csv").name
    content = await file.read(20 * 1024 * 1024 + 1)
    if len(content) > 20 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="分镜表请控制在 20MB 以内")
    try:
        rows = _extract_storyboard_rows(filename, content)
        targets = [item.strip() for item in prompt_targets.split(",") if item.strip() in {"h3", "seedance"}]
        shots = await project_service.import_storyboard(
            project_id,
            rows,
            user_suggestions,
            targets,
            h3_skill_id,
        )
        return {"shots": shots, "project": store.get_project(project_id), "imported_rows": len(rows)}
    except KeyError as exc:
        raise _not_found(str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        logger.exception("项目 %s 导入分镜表失败", project_id)
        raise HTTPException(status_code=500, detail=f"导入分镜表失败: {exc}") from exc


@router.post("/{project_id}/script/upload")
async def upload_script(project_id: str, file: UploadFile = File(...)):
    project = project_service.require_project(project_id)
    filename = Path(file.filename or "script.txt").name
    content = await file.read(10 * 1024 * 1024 + 1)
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="剧本文件请控制在 10MB 以内")
    try:
        text = _extract_script_text(filename, content)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if not text:
        raise HTTPException(status_code=400, detail="文件中没有提取到可用文字")
    destination_dir = project_service.project_dir(project_id) / "documents"
    destination_dir.mkdir(parents=True, exist_ok=True)
    destination = destination_dir / f"{uuid4().hex}{Path(filename).suffix.lower()}"
    destination.write_bytes(content)
    asset = project_service.register_existing_asset(
        project_id,
        destination,
        role=AssetRole.OTHER,
        name=filename,
        description="客户上传的原始剧情脚本",
    )
    asset.mime_type = file.content_type or mimetypes.guess_type(filename)[0] or "application/octet-stream"
    store.save_asset(asset)
    project.brief.story = text[:300_000]
    project.ai_recommended_shot_count = None
    project_service.update_project(project)
    return {"asset": asset, "project": project, "extracted_characters": len(project.brief.story)}


@router.post("/{project_id}/storyboard/approve")
async def approve_storyboard(project_id: str, request: StoryboardApprovalRequest):
    try:
        project = project_service.approve_storyboard(project_id, request.approved, request.comment)
        review = Review(
            project_id=project_id,
            target_type="storyboard",
            target_id=project_id,
            decision=ApprovalStatus.APPROVED if request.approved else ApprovalStatus.CHANGES_REQUESTED,
            comment=request.comment,
            reviewer=request.reviewer,
        )
        store.save_review(review)
        return {"project": project, "review": review}
    except KeyError as exc:
        raise _not_found(str(exc)) from exc


@router.post("/{project_id}/storyboard/reorder")
async def reorder_storyboard(project_id: str, request: ReorderRequest):
    shots = {shot.id: shot for shot in store.list_shots(project_id)}
    if set(request.shot_ids) != set(shots):
        raise HTTPException(status_code=400, detail="Reorder list must contain every shot exactly once")
    ordered = []
    for ordinal, shot_id in enumerate(request.shot_ids, start=1):
        shot = shots[shot_id]
        shot.ordinal = ordinal
        ordered.append(store.save_shot(shot))
    project_service.invalidate_storyboard(project_id)
    return ordered


@router.get("/{project_id}/storyboard.csv")
async def export_storyboard(project_id: str):
    project = store.get_project(project_id)
    if project is None:
        raise _not_found("Project not found")
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["镜号", "时长", "叙事", "对白", "场景", "景别", "机位", "运镜", "主体动作", "通用分镜提示词", "首帧提示词", "H3提示词", "生成模式", "审核状态"])
    storyboard_changed = False
    for shot in store.list_shots(project_id):
        writer.writerow(
            [
                shot.ordinal,
                shot.duration_seconds,
                shot.narrative,
                shot.dialogue,
                shot.scene_description,
                shot.shot_size,
                shot.camera_angle,
                shot.camera_motion,
                shot.subject_motion,
                shot.visual_prompt,
                shot.keyframe_prompt,
                shot.video_prompt,
                (shot.resolved_generation_mode or shot.generation_mode).value,
                shot.approval_status.value,
            ]
        )
    payload = "\ufeff" + output.getvalue()
    return StreamingResponse(
        iter([payload.encode("utf-8")]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{project.id}-storyboard.csv"'},
    )


@router.put("/{project_id}/shots/{shot_id}")
async def update_shot(project_id: str, shot_id: str, shot: Shot):
    if shot.id != shot_id or shot.project_id != project_id:
        raise HTTPException(status_code=400, detail="Shot identity mismatch")
    if store.get_shot(shot_id) is None:
        raise _not_found("Shot not found")
    try:
        return project_service.update_shot(project_id, shot)
    except ShotVersionConflictError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.get("/{project_id}/shots/{shot_id}/h3-prompt-history")
async def h3_prompt_history(project_id: str, shot_id: str):
    shot = store.get_shot(shot_id)
    if shot is None or shot.project_id != project_id:
        raise _not_found("Shot not found")
    return store.list_h3_prompt_history(project_id, shot_id)


@router.patch("/{project_id}/shots/{shot_id}/keyframe-prompt")
async def update_keyframe_prompt(project_id: str, shot_id: str, request: KeyframePromptUpdateRequest):
    try:
        return project_service.update_keyframe_prompt(
            project_id,
            shot_id,
            request.keyframe_prompt,
            request.keyframe_revision_suggestion_draft,
            request.keyframe_revision_mode,
            request.keyframe_reference_asset_ids,
            "keyframe_reference_asset_ids" in request.model_fields_set,
        )
    except KeyError as exc:
        raise _not_found(str(exc)) from exc
    except KeyframeBusyError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.post("/{project_id}/shots/{shot_id}/ai-revise")
async def revise_shot(project_id: str, shot_id: str, request: ShotReviseRequest):
    try:
        return await project_service.revise_shot_with_ai(
            project_id,
            shot_id,
            request.user_suggestions,
            request.prompt_targets,
            request.h3_skill_id,
        )
    except KeyError as exc:
        raise _not_found(str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        logger.exception("项目 %s 的镜头 %s 单镜重做失败", project_id, shot_id)
        raise HTTPException(status_code=500, detail=f"单镜 AI 重做失败: {exc}") from exc


@router.post("/{project_id}/shots/{shot_id}/split-preview")
async def preview_shot_split(project_id: str, shot_id: str, request: ShotSplitPreviewRequest):
    try:
        return await project_service.preview_shot_split(
            project_id,
            shot_id,
            request.user_suggestions,
            request.segment_count,
        )
    except KeyError as exc:
        raise _not_found(str(exc)) from exc
    except ShotVersionConflictError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        logger.exception("项目 %s 的镜头 %s 生成拆分预览失败", project_id, shot_id)
        raise HTTPException(status_code=500, detail=f"生成拆分预览失败: {exc}") from exc


@router.post("/{project_id}/shots/{shot_id}/split-confirm")
async def confirm_shot_split(project_id: str, shot_id: str, request: ShotSplitConfirmRequest):
    try:
        return await project_service.confirm_shot_split(
            project_id,
            shot_id,
            request.preview,
            request.prompt_targets,
            request.h3_skill_id,
        )
    except KeyError as exc:
        raise _not_found(str(exc)) from exc
    except (ShotVersionConflictError, ShotSplitBusyError) as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        logger.exception("项目 %s 的镜头 %s 确认拆分失败", project_id, shot_id)
        raise HTTPException(status_code=500, detail=f"确认拆分失败: {exc}") from exc


@router.get("/{project_id}/shots/{shot_id}/seedance-materials")
async def seedance_materials(project_id: str, shot_id: str):
    try:
        project = project_service.require_project(project_id)
        shot = project_service.require_shot(shot_id)
        if shot.project_id != project_id:
            raise KeyError("Shot not found")
        return project_service.seedance_material_diagnostics(
            project,
            shot,
            store.list_assets(project_id),
        )
    except KeyError as exc:
        raise _not_found(str(exc)) from exc


@router.get("/{project_id}/shots/{shot_id}/keyframe-materials")
async def keyframe_materials(project_id: str, shot_id: str):
    try:
        project = project_service.require_project(project_id)
        shot = project_service.require_shot(shot_id)
        if shot.project_id != project_id:
            raise KeyError("Shot not found")
        return project_service.keyframe_material_diagnostics(
            project,
            shot,
            store.list_assets(project_id),
        )
    except KeyError as exc:
        raise _not_found(str(exc)) from exc


@router.post("/{project_id}/shots/insert-ai")
async def insert_shot_with_ai(project_id: str, request: ShotInsertRequest):
    try:
        return await project_service.insert_shot_with_ai(
            project_id,
            request.after_shot_id,
            request.user_suggestions,
            request.prompt_targets,
            request.h3_skill_id,
        )
    except KeyError as exc:
        raise _not_found(str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        logger.exception("项目 %s 单独新增分镜失败", project_id)
        raise HTTPException(status_code=500, detail=f"单独新增分镜失败: {exc}") from exc


@router.post("/{project_id}/shots")
async def create_shot(project_id: str, shot: Shot):
    if shot.project_id != project_id:
        raise HTTPException(status_code=400, detail="Shot project mismatch")
    project_service.require_project(project_id)
    if not shot.video_prompt_source:
        shot.video_prompt_source = shot.video_prompt
    saved = store.save_shot(shot)
    project_service.invalidate_storyboard(project_id)
    return saved


@router.post("/{project_id}/shots/blank")
async def create_blank_shot(project_id: str):
    project_service.require_project(project_id)
    shots = store.list_shots(project_id)
    shot = Shot(
        project_id=project_id,
        ordinal=len(shots) + 1,
        title=f"镜头 {len(shots) + 1}",
        h3_turbo=False,
        h3_steps=25,
    )
    store.save_shot(shot)
    project_service.invalidate_storyboard(project_id)
    return shot


@router.delete("/{project_id}/shots/{shot_id}")
async def delete_shot(project_id: str, shot_id: str):
    shots = [shot for shot in store.list_shots(project_id) if shot.id != shot_id]
    if len(shots) == len(store.list_shots(project_id)):
        raise _not_found("Shot not found")
    for ordinal, shot in enumerate(shots, start=1):
        shot.ordinal = ordinal
    store.replace_shots(project_id, shots)
    project_service.invalidate_storyboard(project_id)
    return {"deleted": True}


@router.post("/{project_id}/assets")
async def upload_asset(
    project_id: str,
    file: UploadFile = File(...),
    role: AssetRole = Form(AssetRole.OTHER),
    name: str = Form(""),
    character_id: str | None = Form(None),
    description: str = Form(""),
):
    project_service.require_project(project_id)
    filename = Path(file.filename or "asset.bin").name
    suffix = Path(filename).suffix
    destination_dir = project_service.project_dir(project_id) / "assets"
    destination_dir.mkdir(parents=True, exist_ok=True)
    destination = destination_dir / f"{uuid4().hex}{suffix}"
    with destination.open("wb") as handle:
        shutil.copyfileobj(file.file, handle)
    asset = project_service.register_existing_asset(
        project_id,
        destination,
        role=role,
        name=name or filename,
        description=description,
    )
    asset.mime_type = file.content_type or mimetypes.guess_type(filename)[0] or "application/octet-stream"
    asset.character_id = character_id
    return store.save_asset(asset)


@router.delete("/{project_id}/assets/{asset_id}")
async def delete_asset(project_id: str, asset_id: str):
    asset = store.get_asset(asset_id)
    if asset is None or asset.project_id != project_id:
        raise _not_found("Asset not found")
    if not store.delete_asset(asset_id):
        raise _not_found("Asset not found")
    storyboard_changed = False
    changed_shot_ids: list[str] = []
    for shot in store.list_shots(project_id):
        changed = False
        if asset_id in shot.reference_asset_ids:
            shot.reference_asset_ids = [item for item in shot.reference_asset_ids if item != asset_id]
            changed = True
        if shot.keyframe_reference_asset_ids is not None and asset_id in shot.keyframe_reference_asset_ids:
            shot.keyframe_reference_asset_ids = [item for item in shot.keyframe_reference_asset_ids if item != asset_id]
            changed = True
        if shot.video_reference_asset_ids is not None and asset_id in shot.video_reference_asset_ids:
            shot.video_reference_asset_ids = [item for item in shot.video_reference_asset_ids if item != asset_id]
            changed = True
        if shot.keyframe_asset_id == asset_id:
            shot.keyframe_asset_id = None
            shot.image_path = None
            shot.image_status = "pending"
            changed = True
        if shot.last_frame_asset_id == asset_id:
            shot.last_frame_asset_id = None
            changed = True
        if changed:
            store.save_shot(shot)
            changed_shot_ids.append(shot.id)
            storyboard_changed = True
    project = project_service.require_project(project_id)
    character_changed = False
    for character in project.characters:
        if asset_id in character.reference_asset_ids:
            character.reference_asset_ids = [item for item in character.reference_asset_ids if item != asset_id]
            character_changed = True
    if project.style_profile and asset_id in project.style_profile.reference_asset_ids:
        project.style_profile.reference_asset_ids = [
            item for item in project.style_profile.reference_asset_ids if item != asset_id
        ]
        character_changed = True
    for profile in project.scene_profiles:
        if asset_id in profile.reference_asset_ids:
            profile.reference_asset_ids = [item for item in profile.reference_asset_ids if item != asset_id]
            character_changed = True
    if character_changed:
        store.save_project(project)
        # Recompile local manifests; retain paid H3 prose with an older source
        # revision so reference changes are visible without destroying the text.
        project_service.refresh_shot_prompt_caches(project_id)
        storyboard_changed = True
    elif changed_shot_ids:
        project_service.refresh_shot_prompt_caches(project_id, changed_shot_ids)
    path = resolve_media_path(asset.path).resolve()
    project_dir = project_service.project_dir(project_id).resolve()
    if path.exists() and project_dir in path.parents:
        path.unlink()
    if storyboard_changed:
        project_service.invalidate_storyboard(project_id)
    return {"deleted": True}


@router.post("/{project_id}/keyframes/generate")
async def generate_keyframes(project_id: str, request: KeyframeRequest):
    try:
        return await project_service.generate_keyframes(
            project_id,
            request.shot_ids,
            request.image_provider,
            request.image_model,
            request.revision_mode,
            request.user_suggestions,
        )
    except KeyframeBusyError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        logger.exception("项目 %s 的分镜首帧生成失败", project_id)
        raise HTTPException(status_code=500, detail=str(exc)) from exc


@router.post("/{project_id}/shots/{shot_id}/keyframe/upload")
async def upload_keyframe(project_id: str, shot_id: str, file: UploadFile = File(...)):
    shot = store.get_shot(shot_id)
    if shot is None or shot.project_id != project_id:
        raise _not_found("Shot not found")
    filename = Path(file.filename or "keyframe.png").name
    suffix = Path(filename).suffix.lower()
    if suffix not in {".png", ".jpg", ".jpeg", ".webp", ".bmp"}:
        raise HTTPException(status_code=400, detail="仅支持 PNG/JPG/WebP/BMP 格式的首帧图片")
    content = await file.read(20 * 1024 * 1024 + 1)
    if len(content) > 20 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="首帧图片请控制在 20MB 以内")
    # File reads yield to prompt/image tasks. Merge the upload into current data.
    shot = store.get_shot(shot_id)
    if shot is None or shot.project_id != project_id:
        raise _not_found("Shot not found")
    try:
        project_service.ensure_keyframe_idle(shot_id)
    except KeyframeBusyError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    image_dir = project_service.project_dir(project_id) / "images"
    image_dir.mkdir(parents=True, exist_ok=True)
    destination = image_dir / f"{shot.id}-upload-{uuid4().hex}{suffix}"
    destination.write_bytes(content)
    asset = project_service.register_existing_asset(
        project_id,
        destination,
        role=AssetRole.KEYFRAME,
        name=f"镜头 {shot.ordinal} 首帧（本地上传）",
        description=shot.keyframe_prompt or "",
    )
    shot.keyframe_asset_id = asset.id
    shot.image_path = str(destination)
    shot.image_status = "completed"
    shot.keyframe_prompt_source_revision = shot.content_revision
    shot.version += 1
    return store.save_shot(shot)


@router.post("/{project_id}/h3-prompts/generate")
async def generate_h3_prompts(project_id: str, request: H3PromptGenerateRequest):
    try:
        return await project_service.generate_h3_prompts(
            project_id,
            request.shot_ids,
            request.skill_id,
            request.user_suggestions,
        )
    except KeyError as exc:
        raise _not_found(str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        logger.exception("项目 %s 的 H3 Prompt Skill 生成失败", project_id)
        raise HTTPException(status_code=500, detail=f"H3 Prompt 生成失败: {exc}") from exc


@router.post("/{project_id}/seedance-prompts/generate")
async def generate_seedance_prompts(project_id: str, request: SeedancePromptGenerateRequest):
    try:
        return project_service.generate_seedance_prompts(project_id, request.shot_ids)
    except KeyError as exc:
        raise _not_found(str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.post("/{project_id}/seedance/estimate")
async def estimate_seedance(project_id: str, request: SeedanceEstimateRequest):
    try:
        project = project_service.require_project(project_id)
        shots = store.list_shots(project_id)
        if request.shot_ids is not None:
            selected = set(request.shot_ids)
            shots = [shot for shot in shots if shot.id in selected]
        if not shots:
            raise ValueError("请至少选择一个镜头")
        model = resolve_seedance_model(request.model_id)
        resolution = validate_seedance_resolution(model.id, request.resolution)
        assets = store.list_assets(project_id)
        estimates = []
        for shot in shots:
            refs = project_service.seedance_reference_assets(shot, assets)
            estimates.append(
                estimate_seedance_cost(
                    model.id,
                    resolution,
                    project.brief.aspect_ratio,
                    [shot.duration_seconds],
                    has_video_input=any(asset.type == AssetType.VIDEO for asset in refs),
                )
            )
        total_tokens = sum(int(item["estimated_tokens"]) for item in estimates)
        total_cost = sum(float(item["estimated_yuan"]) for item in estimates)
        total_billed = sum(int(item["billed_duration_seconds"]) for item in estimates)
        result = dict(estimates[0])
        result.update(
            {
                "shot_count": len(estimates),
                "requested_duration_seconds": round(sum(shot.duration_seconds for shot in shots), 3),
                "billed_duration_seconds": total_billed,
                "estimated_tokens": total_tokens,
                "unit_price_per_million_tokens": round(total_cost * 1_000_000 / total_tokens, 4) if total_tokens else 0,
                "estimated_yuan": round(total_cost, 4),
                "estimated_yuan_per_second": round(
                    sum(float(item["estimated_yuan_per_second"]) * int(item["billed_duration_seconds"]) for item in estimates) / total_billed,
                    4,
                ) if total_billed else 0,
                "estimated_yuan_per_task": round(sum(float(item["estimated_yuan_per_task"]) for item in estimates) / len(estimates), 4),
                "has_video_input": any(bool(item["has_video_input"]) for item in estimates),
                "note": "已逐镜按是否包含参考视频使用对应单价估算；最终以方舟任务 usage 与账单为准。",
            }
        )
        return result
    except (KeyError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.get("/{project_id}/seedance/preflight")
async def seedance_preflight(project_id: str, model_id: str | None = None, resolution: str | None = None):
    project_service.require_project(project_id)
    try:
        return await SeedanceClient().preflight(model_id, resolution)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


def _safe_archive_component(value: str, fallback: str) -> str:
    clean = re.sub(r'[\x00-\x1f<>:"/\\|?*]+', "_", (value or "").strip()).strip(" ._")
    return (clean or fallback)[:80]


def _build_keyframe_archive(project_id: str, shot_ids: list[str] | None) -> tuple[Path, str, int]:
    project = project_service.require_project(project_id)
    selected = set(shot_ids) if shot_ids is not None else None
    entries: list[tuple[Shot, Path]] = []
    for shot in store.list_shots(project_id):
        if selected is not None and shot.id not in selected:
            continue
        asset = store.get_asset(shot.keyframe_asset_id) if shot.keyframe_asset_id else None
        if asset is None or asset.project_id != project_id:
            continue
        path = resolve_media_path(asset.path)
        if path.is_file():
            entries.append((shot, path))
    if not entries:
        raise ValueError("所选镜头中没有可导出的已生成分镜图")

    handle = tempfile.NamedTemporaryFile(prefix="videoworkflow-keyframes-", suffix=".zip", delete=False)
    archive_path = Path(handle.name)
    handle.close()
    try:
        with zipfile.ZipFile(archive_path, "w", compression=zipfile.ZIP_STORED, allowZip64=True) as archive:
            for shot, source in entries:
                title = _safe_archive_component(shot.title, f"镜头_{shot.ordinal:03d}")
                suffix = source.suffix.lower() or ".png"
                archive.write(source, arcname=f"镜头_{shot.ordinal:03d}_{title}{suffix}")
    except Exception:
        archive_path.unlink(missing_ok=True)
        raise
    project_title = _safe_archive_component(project.brief.title, project.id)
    return archive_path, f"{project_title}_分镜图_{len(entries)}张.zip", len(entries)


@router.post("/{project_id}/keyframes/export")
async def export_keyframes(project_id: str, request: KeyframeExportRequest):
    try:
        archive_path, filename, _ = await asyncio.to_thread(
            _build_keyframe_archive,
            project_id,
            request.shot_ids,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return FileResponse(
        archive_path,
        media_type="application/zip",
        filename=filename,
        background=BackgroundTask(archive_path.unlink, missing_ok=True),
    )


@router.post("/{project_id}/render/plan")
async def plan_render(project_id: str):
    try:
        shots = project_service.plan_all_shots(project_id)
        project = project_service.require_project(project_id)
        project.status = ProjectStatus.RENDER_PLAN_APPROVED
        store.save_project(project)
        return shots
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.post("/{project_id}/render")
async def enqueue_render(project_id: str, request: RenderRequest):
    try:
        if request.provider == "ark_seedance":
            return render_queue.enqueue_seedance(
                project_id,
                request.shot_ids,
                model_id=request.model_id,
                resolution=request.resolution,
                generate_audio=request.generate_audio,
            )
        return render_queue.enqueue(project_id, request.shot_ids)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.get("/{project_id}/jobs")
async def list_jobs(project_id: str):
    project_service.require_project(project_id)
    return store.list_jobs(project_id)


@router.post("/{project_id}/jobs/{job_id}/cancel")
async def cancel_job(project_id: str, job_id: str):
    job = store.get_job(job_id)
    if job is None or job.project_id != project_id:
        raise _not_found("Job not found")
    return await render_queue.cancel(job_id)


@router.post("/{project_id}/jobs/{job_id}/select")
async def select_render_job(project_id: str, job_id: str):
    try:
        return await asyncio.to_thread(render_queue.select_output, project_id, job_id)
    except KeyError as exc:
        raise _not_found(str(exc)) from exc
    except (ValueError, OSError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.delete("/{project_id}/jobs/{job_id}")
async def delete_job(project_id: str, job_id: str):
    job = store.get_job(job_id)
    if job is None or job.project_id != project_id:
        raise _not_found("Job not found")
    if job.status not in {JobStatus.COMPLETED, JobStatus.FAILED, JobStatus.CANCELLED}:
        raise HTTPException(status_code=409, detail="任务仍在进行中，请先取消后再删除")
    shot = store.get_shot(job.shot_id) if job.shot_id else None
    if shot and shot.selected_video_job_id == job.id:
        raise HTTPException(status_code=409, detail="该任务是当前采用版本，请先采用另一版本再删除")
    was_current_output = bool(shot and job.output_path and shot.video_path == job.output_path)
    store.delete_job(job_id)
    if shot:
        remaining = [item for item in store.list_jobs(project_id) if item.shot_id == job.shot_id]
        if was_current_output:
            fallback = next(
                (
                    item for item in remaining
                    if item.status == JobStatus.COMPLETED and item.output_path
                ),
                None,
            )
            if fallback:
                shot.video_path = fallback.output_path
                shot.last_frame_asset_id = render_queue.job_last_frame_asset_id(fallback)
                shot.video_status = "completed"
            else:
                shot.video_path = None
                shot.last_frame_asset_id = None
                latest = remaining[0] if remaining else None
                shot.video_status = (
                    "failed" if latest and latest.status == JobStatus.FAILED
                    else latest.status.value if latest
                    else "pending"
                )
        elif shot.selected_video_job_id and shot.video_path:
            shot.video_status = "completed"
        elif not shot.video_path:
            latest = remaining[0] if remaining else None
            shot.video_status = (
                "failed" if latest and latest.status == JobStatus.FAILED
                else latest.status.value if latest
                else "pending"
            )
        store.save_shot(shot)
    return {"deleted": True}


@router.get("/{project_id}/comfyui/preflight")
async def comfyui_preflight(project_id: str):
    project_service.require_project(project_id)
    if settings.H3_PROVIDER == "metaso_h3":
        try:
            return await MetaSoH3Client().preflight()
        except Exception as exc:
            return {"online": False, "ok": False, "provider": "metaso_h3", "error": str(exc)}
    if settings.H3_PROVIDER == "atlas_h3":
        return {
            "online": bool(settings.ATLASCLOUD_API_KEY),
            "ok": bool(settings.ATLASCLOUD_API_KEY),
            "provider": "atlas_h3",
            "model": "minimax/h3-developer/reference-to-video",
            "resolution": settings.H3_ATLAS_RESOLUTION,
            "ratio": settings.H3_ATLAS_RATIO,
            "message": "Atlas H3 配置已就绪" if settings.ATLASCLOUD_API_KEY else "Atlas H3 API Key 尚未配置",
        }
    client = ComfyUIClient()
    builder = H3WorkflowBuilder()
    try:
        system = await client.health()
        model_profile = resolve_h3_model_profile(None)
        text_encoder_profile = resolve_h3_text_encoder_profile(None)
        selected_models = [
            h3_diffusion_model(model_profile, GenerationMode.I2V),
            h3_diffusion_model(model_profile, GenerationMode.R2V),
            h3_text_encoder(text_encoder_profile),
        ]
        preflight = await client.preflight(builder.preflight_requirements(selected_models))
        return {
            "online": True,
            "system": system,
            "model_profile": model_profile,
            "text_encoder_profile": text_encoder_profile,
            **preflight,
        }
    except Exception as exc:
        return {"online": False, "ok": False, "error": str(exc)}


@router.post("/{project_id}/subtitles/generate")
async def generate_subtitles(project_id: str):
    project_service.require_project(project_id)
    lines: list[str] = []
    cursor = 0.0
    for index, shot in enumerate(store.list_shots(project_id), start=1):
        text = (shot.dialogue or shot.narrative).strip()
        if not text:
            cursor += shot.duration_seconds
            continue
        start = cursor
        end = cursor + shot.duration_seconds
        lines.extend([str(index), f"{_srt_time(start)} --> {_srt_time(end)}", text, ""])
        cursor = end
    path = project_service.project_dir(project_id) / "assets" / f"subtitles-{uuid4().hex}.srt"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")
    return project_service.register_existing_asset(project_id, path, AssetRole.OTHER, "自动字幕")


@router.post("/{project_id}/finalize")
async def finalize_project(project_id: str, request: FinalizeRequest):
    try:
        return await asyncio.to_thread(finalizer.finalize, project_id, request)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.get("/{project_id}/download/{kind}/{item_id}")
async def download_project_file(project_id: str, kind: str, item_id: str, inline: bool = False):
    if kind == "asset":
        item = store.get_asset(item_id)
        path = resolve_media_path(item.path) if item and item.project_id == project_id else None
    elif kind == "job":
        item = store.get_job(item_id)
        path = resolve_media_path(item.output_path) if item and item.project_id == project_id and item.output_path else None
    elif kind == "delivery":
        item = next((delivery for delivery in store.list_deliveries(project_id) if delivery.id == item_id), None)
        path = resolve_media_path(item.output_path) if item else None
    elif kind == "preview":
        item = next((delivery for delivery in store.list_deliveries(project_id) if delivery.id == item_id), None)
        path = resolve_media_path(item.preview_path) if item and item.preview_path else None
    else:
        raise HTTPException(status_code=400, detail="Invalid file kind")
    if path is None or not path.exists():
        raise _not_found("File not found")
    media_type = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    if inline:
        return FileResponse(path, media_type=media_type)
    return FileResponse(path, filename=path.name, media_type=media_type)


@public_router.get("/{token}")
async def public_review(token: str):
    project = next((item for item in store.list_projects() if item.review_token == token), None)
    if project is None:
        raise _not_found("Review link not found")
    bundle = _bundle(project.id)
    bundle["project"] = project.model_copy(update={"review_token": ""})
    bundle["assets"] = [asset.model_copy(update={"path": ""}) for asset in bundle["assets"]]
    bundle["jobs"] = [job.model_copy(update={"workflow_snapshot": {}, "input_snapshot": {}, "output_path": None}) for job in bundle["jobs"]]
    bundle["deliveries"] = [
        delivery.model_copy(
            update={
                "output_path": "",
                "preview_path": "available" if delivery.preview_path else None,
                "subtitle_path": None,
            }
        )
        for delivery in bundle["deliveries"]
    ]
    return bundle


@public_router.post("/{token}")
async def submit_public_review(token: str, request: PublicReviewRequest):
    project = next((item for item in store.list_projects() if item.review_token == token), None)
    if project is None:
        raise _not_found("Review link not found")
    if request.decision not in {ApprovalStatus.APPROVED, ApprovalStatus.CHANGES_REQUESTED}:
        raise HTTPException(status_code=400, detail="Review decision must be approved or changes_requested")
    if request.target_type == "shot":
        shot = store.get_shot(request.target_id)
        if shot is None or shot.project_id != project.id:
            raise _not_found("Shot not found")
    elif request.target_type == "storyboard":
        if request.target_id != project.id:
            raise HTTPException(status_code=400, detail="Storyboard target mismatch")
    elif request.target_type == "delivery":
        delivery = next((item for item in store.list_deliveries(project.id) if item.id == request.target_id), None)
        if delivery is None:
            raise _not_found("Delivery not found")
    else:
        raise HTTPException(status_code=400, detail="Unsupported review target")
    review = Review(
        project_id=project.id,
        target_type=request.target_type,
        target_id=request.target_id,
        decision=request.decision,
        comment=request.comment,
        reviewer=request.reviewer,
    )
    store.save_review(review)
    if request.target_type == "shot":
        shot.approval_status = request.decision
        store.save_shot(shot)
    elif request.target_type == "storyboard":
        project.status = ProjectStatus.STORYBOARD_APPROVED if request.decision == ApprovalStatus.APPROVED else ProjectStatus.STORYBOARD_DRAFT
        store.save_project(project)
        for shot in store.list_shots(project.id):
            shot.approval_status = request.decision
            store.save_shot(shot)
    elif request.target_type == "delivery":
        project.status = ProjectStatus.DELIVERED if request.decision == ApprovalStatus.APPROVED else ProjectStatus.FINAL_REVIEW
        store.save_project(project)
    return review


def _srt_time(seconds: float) -> str:
    millis = int(round(seconds * 1000))
    hours, remainder = divmod(millis, 3_600_000)
    minutes, remainder = divmod(remainder, 60_000)
    secs, ms = divmod(remainder, 1000)
    return f"{hours:02d}:{minutes:02d}:{secs:02d},{ms:03d}"
